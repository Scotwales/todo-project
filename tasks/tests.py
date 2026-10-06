import io
import sqlite3
import tempfile
from contextlib import closing
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, connections
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from django.db.migrations.executor import MigrationExecutor
from openpyxl import load_workbook
from .models import (
    DesktopNotification, PomodoroSession, RecurringSchedule, Task, TaskCompletion,
    TaskExtensionHistory, TaskTimeEntry,
)
from .planning import generate_occurrences, tracked_seconds_between


class TaskModelTests(TestCase):
    def test_recurrence_occurrence_dates(self):
        task = Task.objects.create(title="Weekly", due_date=date(2025, 1, 6), recurrence=Task.Recurrence.WEEKLY)
        self.assertTrue(task.occurs_on(date(2025, 1, 13)))
        self.assertFalse(task.occurs_on(date(2025, 1, 14)))

    def test_weekly_task_without_due_date_uses_creation_weekday(self):
        task = Task.objects.create(title="Weekly", recurrence=Task.Recurrence.WEEKLY)
        created_day = timezone.localtime(task.created_at).date()
        self.assertTrue(task.occurs_on(created_day))
        self.assertFalse(task.occurs_on(created_day + timedelta(days=1)))

    def test_monthly_recurrence_clamps_to_month_end(self):
        task = Task.objects.create(title="Monthly", due_date=date(2025, 1, 31), recurrence=Task.Recurrence.MONTHLY)
        self.assertTrue(task.occurs_on(date(2025, 2, 28)))
        self.assertFalse(task.occurs_on(date(2025, 2, 27)))

    def test_weekdays_and_custom_recurring_schedule(self):
        today = timezone.localdate()
        weekday_task = Task.objects.create(title="Weekdays", due_date=today, recurrence=Task.Recurrence.WEEKDAYS)
        self.assertEqual(weekday_task.occurs_on(today), today.weekday() < 5)
        task = Task.objects.create(title="Custom", due_date=today, recurrence=Task.Recurrence.CUSTOM)
        schedule = RecurringSchedule.objects.create(
            template=task, frequency=RecurringSchedule.Frequency.CUSTOM, starts_on=today,
            weekdays=[0, 2, 4],
        )
        for offset in range(7):
            day = today + timedelta(days=offset)
            self.assertEqual(schedule.occurs_on(day), day.weekday() in [0, 2, 4])

    def test_task_completion_is_unique_for_each_day(self):
        task = Task.objects.create(title="Daily", recurrence=Task.Recurrence.DAILY)
        day = date(2025, 1, 1)
        TaskCompletion.objects.create(task=task, occurrence_date=day)
        with self.assertRaises(IntegrityError):
            TaskCompletion.objects.create(task=task, occurrence_date=day)


class TaskViewTests(TransactionTestCase):
    def test_about_screen_displays_product_version_build_date_and_release_notes(self):
        response = self.client.get(reverse("about"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "About Daily")
        self.assertContains(response, "1.1.0")
        self.assertContains(response, "Build date")
        self.assertContains(response, "Release notes")

    def test_theme_preference_is_persisted_in_local_settings_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            preferences = Path(temporary) / "settings" / "preferences.json"
            with override_settings(PREFERENCES_FILE=preferences):
                response = self.client.post(reverse("settings"), {"theme": "dark"})
                self.assertEqual(response.status_code, 302)
                self.assertEqual(
                    self.client.get(reverse("settings")).context["saved_theme"], "dark"
                )
            self.assertIn('"theme": "dark"', preferences.read_text(encoding="utf-8"))


    def test_dashboard_shows_today_and_quick_add(self):
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Quick add")

    def test_bulk_creation_applies_shared_task_settings(self):
        response = self.client.post(reverse("bulk_create"), {
            "titles": "Write proposal\nReview roadmap\n\nSend update",
            "category": "Work",
            "priority": str(Task.Priority.HIGH),
            "due_date": timezone.localdate().isoformat(),
            "due_time": "15:00",
            "reminder_interval_minutes": "15",
            "deliverable_minutes": "45",
        })
        self.assertEqual(response.status_code, 302)
        tasks = list(Task.objects.filter(category="Work"))
        self.assertEqual(len(tasks), 3)
        self.assertTrue(all(task.priority == Task.Priority.HIGH for task in tasks))
        self.assertTrue(all(task.due_time.isoformat() == "15:00:00" for task in tasks))
        self.assertTrue(all(task.deliverable_minutes == 45 for task in tasks))
        self.assertTrue(all(task.reminder_interval_minutes == 15 for task in tasks))
        expected_reminder = timezone.make_aware(
            datetime.combine(timezone.localdate(), datetime.strptime("15:00", "%H:%M").time()),
            timezone.get_current_timezone(),
        ) - timedelta(minutes=15)
        self.assertTrue(all(task.reminder_at == expected_reminder for task in tasks))

    def test_bulk_creation_requires_due_time_for_reminder_interval(self):
        response = self.client.post(reverse("bulk_create"), {
            "titles": "Write proposal", "reminder_interval_minutes": "15",
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Task.objects.exists())
        self.assertContains(response, "Add a due time to schedule an interval reminder.")

    def test_recurring_generation_creates_unique_child_occurrences(self):
        today = timezone.localdate()
        template = Task.objects.create(
            title="Daily review", due_date=today, due_time="09:00",
            recurrence=Task.Recurrence.DAILY, deliverable_minutes=20,
        )
        schedule = RecurringSchedule.objects.create(
            template=template, frequency=RecurringSchedule.Frequency.DAILY,
            starts_on=today, last_generated_through=today - timedelta(days=1),
        )
        through = today + timedelta(days=3)
        self.assertEqual(generate_occurrences(through=through), 4)
        self.assertEqual(generate_occurrences(through=through), 0)
        self.assertEqual(Task.objects.filter(recurring_template=template).count(), 4)
        self.assertEqual(Task.objects.filter(recurrence__in=("daily", "weekdays", "weekly", "monthly", "custom")).count(), 1)
        self.assertEqual(Task.objects.filter(recurring_template=template, is_completed=False).first().deliverable_minutes, 20)
        schedule.refresh_from_db()
        self.assertEqual(schedule.last_generated_through, through)

    def test_recurring_schedule_migration_and_templates_do_not_duplicate(self):
        today = timezone.localdate()
        template = Task.objects.create(
            title="Migration review", due_date=today, recurrence=Task.Recurrence.DAILY,
        )
        schedule = RecurringSchedule.objects.create(
            template=template, frequency=RecurringSchedule.Frequency.DAILY, starts_on=today,
        )
        child_date = today + timedelta(days=1)
        child = Task.objects.create(
            title=template.title, due_date=child_date, recurring_template=template,
            occurrence_date=child_date,
        )
        self.assertEqual(generate_occurrences(through=child_date), 1)
        child.refresh_from_db()
        self.assertFalse(child.is_cancelled)
        self.assertEqual(Task.objects.filter(pk=template.pk).count(), 1)
        self.assertEqual(Task.objects.filter(recurring_template=template).count(), 2)
        self.assertEqual(schedule.template_id, template.pk)

    def test_daily_planner_and_analytics_show_planning_metrics(self):
        today = timezone.localdate()
        task = Task.objects.create(
            title="Plan the sprint", due_date=today, due_time="11:00",
            deliverable_minutes=60, is_completed=True, completed_at=timezone.now(),
        )
        TaskTimeEntry.objects.create(
            task=task, started_at=timezone.now() - timedelta(minutes=30),
            ended_at=timezone.now(), duration_seconds=1800,
        )
        TaskExtensionHistory.objects.create(
            task=task, extension_type=TaskExtensionHistory.ExtensionType.DURATION,
            minutes=15,
        )
        planner_response = self.client.get(reverse("planner"))
        self.assertEqual(planner_response.status_code, 200)
        self.assertContains(planner_response, "Plan the sprint")
        analytics_response = self.client.get(reverse("analytics"))
        self.assertEqual(analytics_response.status_code, 200)
        self.assertEqual(analytics_response.context["completed_count"], 1)
        self.assertEqual(analytics_response.context["extended_count"], 1)
        self.assertEqual(analytics_response.context["estimated_hours"], 1.0)
        self.assertEqual(analytics_response.context["actual_hours"], .5)

    def test_planner_positions_estimates_and_avoids_overlapping_flexible_tasks(self):
        today = timezone.localdate()
        Task.objects.create(title="Design review", due_date=today, due_time="10:00", deliverable_minutes=60)
        Task.objects.create(title="Client call", due_date=today, due_time="10:30", deliverable_minutes=60)
        Task.objects.create(title="Flexible task", due_date=today, deliverable_minutes=30)
        response = self.client.get(reverse("planner"), {"date": today.isoformat()})
        self.assertEqual(response.status_code, 200)
        schedule = response.context["schedule"]
        flexible = next(item for item in schedule if item["task"].title == "Flexible task")
        self.assertEqual(flexible["start"].strftime("%H:%M"), "10:30")
        timed = [item for item in schedule if item["task"].title != "Flexible task"]
        self.assertEqual({item["left"] for item in timed}, {0, 50})

    def test_time_analytics_clips_sessions_at_day_boundaries(self):
        today = timezone.localdate()
        midnight = timezone.make_aware(
            datetime.combine(today, datetime.min.time()), timezone.get_current_timezone()
        )
        entry = TaskTimeEntry(
            task=Task.objects.create(title="Overnight focus"),
            started_at=midnight - timedelta(hours=1),
            ended_at=midnight + timedelta(minutes=30),
        )
        self.assertEqual(
            tracked_seconds_between(midnight, midnight + timedelta(days=1), [entry]),
            1800,
        )

    def test_custom_recurrence_form_persists_selected_weekdays(self):
        response = self.client.post(reverse("create_task"), {
            "title": "Practice piano", "due_date": timezone.localdate().isoformat(),
            "recurrence": Task.Recurrence.CUSTOM, "weekdays": ["0", "2", "4"],
            "priority": Task.Priority.NORMAL,
        })
        self.assertEqual(response.status_code, 302)
        task = Task.objects.get(title="Practice piano")
        self.assertEqual(task.recurring_schedule.weekdays, [0, 2, 4])

    def test_ending_recurring_series_cancels_future_children_without_losing_history(self):
        today = timezone.localdate()
        template = Task.objects.create(
            title="Daily review", due_date=today, recurrence=Task.Recurrence.DAILY,
        )
        RecurringSchedule.objects.create(
            template=template, frequency=RecurringSchedule.Frequency.DAILY, starts_on=today,
        )
        generate_occurrences(through=today + timedelta(days=3))
        occurrence = template.occurrences.filter(occurrence_date__gt=today).first()
        TaskExtensionHistory.objects.create(
            task=occurrence, extension_type=TaskExtensionHistory.ExtensionType.DURATION,
            minutes=15,
        )
        response = self.client.post(reverse("update_task", args=[occurrence.id]), {
            "title": "Daily review", "notes": "", "category": "", "due_date": today.isoformat(),
            "due_time": "", "reminder_at": "", "reminder_interval_minutes": "",
            "deliverable_minutes": "", "recurrence": Task.Recurrence.NONE,
            "priority": Task.Priority.NORMAL, "apply_to_series": "on",
        })
        self.assertEqual(response.status_code, 302)
        occurrence.refresh_from_db()
        self.assertTrue(occurrence.is_cancelled)
        self.assertEqual(TaskExtensionHistory.objects.filter(task=occurrence).count(), 1)
        self.assertFalse(RecurringSchedule.objects.filter(template=template).exists())

    def test_task_extension_options_record_history(self):
        today = timezone.localdate()
        task = Task.objects.create(title="Finish notes", due_date=today, due_time="15:00")
        response = self.client.post(reverse("extend_task", args=[task.id]), {"option": "30"})
        self.assertEqual(response.status_code, 200)
        task.refresh_from_db()
        self.assertEqual(task.due_time.strftime("%H:%M"), "15:30")
        history = TaskExtensionHistory.objects.get(task=task)
        self.assertEqual(history.minutes, 30)
        self.assertEqual(history.extension_type, TaskExtensionHistory.ExtensionType.DURATION)

    def test_custom_and_tomorrow_extensions_are_validated_and_tracked(self):
        task = Task.objects.create(title="Prepare slides", due_date=timezone.localdate(), due_time="10:00")
        invalid = self.client.post(reverse("extend_task", args=[task.id]), {
            "option": "custom", "custom_minutes": "1441",
        })
        self.assertEqual(invalid.status_code, 400)
        custom = self.client.post(reverse("extend_task", args=[task.id]), {
            "option": "custom", "custom_minutes": "75",
        })
        self.assertEqual(custom.status_code, 200)
        history = TaskExtensionHistory.objects.filter(task=task).first()
        self.assertEqual(history.minutes, 75)
        self.assertEqual(history.extension_type, TaskExtensionHistory.ExtensionType.CUSTOM)
        tomorrow = self.client.post(reverse("extend_task", args=[task.id]), {"option": "tomorrow"})
        self.assertEqual(tomorrow.status_code, 200)
        latest_history = TaskExtensionHistory.objects.filter(task=task).first()
        self.assertEqual(latest_history.extension_type, TaskExtensionHistory.ExtensionType.TOMORROW)
        self.assertEqual(TaskExtensionHistory.objects.filter(task=task).count(), 2)

    def test_timer_start_and_pause_track_actual_time(self):
        task = Task.objects.create(title="Draft report")
        started = self.client.post(reverse("task_timer", args=[task.id, "start"]))
        self.assertEqual(started.status_code, 200)
        entry = TaskTimeEntry.objects.get(task=task, ended_at__isnull=True)
        entry.started_at = timezone.now() - timedelta(minutes=2)
        entry.save(update_fields=("started_at",))
        paused = self.client.post(reverse("task_timer", args=[task.id, "pause"]))
        self.assertEqual(paused.status_code, 200)
        entry.refresh_from_db()
        self.assertGreaterEqual(entry.duration_seconds, 120)

    def test_notification_actions_complete_task_or_extend_due_time(self):
        task = Task.objects.create(title="Send invoice", due_date=timezone.localdate(), due_time="10:00")
        item = DesktopNotification.objects.create(task=task)
        response = self.client.post(reverse("notification_action", args=[item.id, "extend"]), {
            "option": "15",
        })
        self.assertEqual(response.status_code, 200)
        task.refresh_from_db()
        self.assertEqual(task.due_time.strftime("%H:%M"), "10:15")
        self.assertEqual(TaskExtensionHistory.objects.filter(task=task).count(), 1)
        second = DesktopNotification.objects.create(task=task)
        response = self.client.post(reverse("notification_action", args=[second.id, "done"]))
        self.assertEqual(response.status_code, 200)
        task.refresh_from_db()
        self.assertTrue(task.is_completed)

    def test_create_and_toggle_task(self):
        today = timezone.localdate().isoformat()
        response = self.client.post(reverse("create_task"), {"title": "Read", "due_date": today, "recurrence": "none", "priority": 2})
        self.assertEqual(response.status_code, 302)
        task = Task.objects.get(title="Read")
        self.client.post(reverse("toggle_task", args=[task.id]), {"date": today})
        task.refresh_from_db()
        self.assertTrue(task.is_completed)

    def test_update_task_can_keep_due_date_empty(self):
        task = Task.objects.create(title="Unscheduled")
        response = self.client.post(reverse("update_task", args=[task.id]), {
            "title": "Updated task", "notes": "Some notes", "due_date": "", "due_time": "",
            "reminder_at": "", "recurrence": "none", "priority": 3,
        })
        self.assertEqual(response.status_code, 302)
        task.refresh_from_db()
        self.assertEqual(task.title, "Updated task")
        self.assertEqual(task.notes, "Some notes")
        self.assertIsNone(task.due_date)
        self.assertEqual(task.priority, Task.Priority.HIGH)

    def test_duplicate_task_copies_planning_details_but_not_completion(self):
        task = Task.objects.create(
            title="Review notes", notes="Keep the key points", category="Work",
            due_date=timezone.localdate(), due_time="14:30", reminder_interval_minutes=15,
            deliverable_minutes=45,
            recurrence=Task.Recurrence.WEEKLY, priority=Task.Priority.HIGH,
            is_completed=True, completed_at=timezone.now(),
        )
        RecurringSchedule.objects.create(
            template=task, frequency=RecurringSchedule.Frequency.WEEKLY,
            weekdays=[], starts_on=task.due_date,
        )
        response = self.client.post(reverse("duplicate_task", args=[task.id]), {"next": "dashboard"})
        self.assertEqual(response.status_code, 302)
        task.refresh_from_db()
        duplicate = Task.objects.exclude(pk=task.id).get()
        self.assertEqual(duplicate.title, task.title)
        self.assertEqual(duplicate.notes, task.notes)
        self.assertEqual(duplicate.due_date, task.due_date)
        self.assertEqual(duplicate.due_time, task.due_time)
        self.assertEqual(duplicate.category, task.category)
        self.assertEqual(duplicate.deliverable_minutes, task.deliverable_minutes)
        self.assertEqual(duplicate.reminder_interval_minutes, task.reminder_interval_minutes)
        self.assertEqual(duplicate.recurrence, task.recurrence)
        self.assertEqual(duplicate.priority, task.priority)
        self.assertEqual(duplicate.recurring_schedule.frequency, task.recurring_schedule.frequency)
        self.assertFalse(duplicate.is_completed)
        self.assertIsNone(duplicate.completed_at)

    def test_task_redirect_does_not_allow_external_host(self):
        response = self.client.post(reverse("create_task"), {
            "title": "Safe redirect", "due_date": timezone.localdate().isoformat(),
            "recurrence": "none", "priority": 2, "next": "https://example.com",
        })
        self.assertEqual(response["Location"], reverse("dashboard"))
        self.assertTrue(Task.objects.filter(title="Safe redirect").exists())

    def test_cannot_toggle_task_on_non_occurrence_date(self):
        task = Task.objects.create(title="Weekly", due_date=timezone.localdate(), recurrence=Task.Recurrence.WEEKLY)
        non_occurrence = timezone.localdate() + timedelta(days=1)
        response = self.client.post(
            reverse("toggle_task", args=[task.id]), {"date": non_occurrence.isoformat()}
        )
        self.assertEqual(response.status_code, 400)

    def test_recurring_toggle_tracks_occurrence_not_entire_task(self):
        task = Task.objects.create(title="Daily", due_date=timezone.localdate(), recurrence=Task.Recurrence.DAILY)
        self.client.post(reverse("toggle_task", args=[task.id]), {"date": timezone.localdate().isoformat()})
        self.assertEqual(TaskCompletion.objects.filter(task=task).count(), 1)
        self.assertFalse(task.is_completed)

    def test_calendar_invalid_month_is_safe(self):
        response = self.client.get(reverse("calendar"), {"year": "9000", "month": "15"})
        self.assertEqual(response.status_code, 200)

    def test_excel_export_contains_task(self):
        Task.objects.create(title="Export me", due_date=timezone.localdate())
        response = self.client.get(reverse("export_excel"))
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.active["A2"].value, "Export me")

    def test_excel_export_keeps_formula_like_titles_as_text(self):
        Task.objects.create(title="=1+1", due_date=timezone.localdate())
        response = self.client.get(reverse("export_excel"))
        cell = load_workbook(io.BytesIO(response.content)).active["A2"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_pomodoro_rejects_invalid_duration(self):
        self.assertEqual(self.client.post(reverse("pomodoro"), {"duration": "999"}).status_code, 400)
        self.client.post(reverse("pomodoro"), {"duration": "25"})
        self.assertEqual(PomodoroSession.objects.filter(completed_at__isnull=True).count(), 1)

    def test_dark_mode_preference_is_saved(self):
        self.client.post(reverse("settings"), {"theme": "dark"})
        self.assertEqual(self.client.session["theme"], "dark")

    def test_backup_download_is_valid_sqlite(self):
        Task.objects.create(title="Backup me", due_date=timezone.localdate())
        response = self.client.get(reverse("backup"))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "backup.sqlite3"
            path.write_bytes(response.content)
            database = sqlite3.connect(str(path))
            try:
                self.assertEqual(database.execute("PRAGMA integrity_check").fetchone()[0], "ok")
                self.assertEqual(database.execute("SELECT title FROM tasks_task").fetchone()[0], "Backup me")
            finally:
                database.close()

    def test_backup_download_is_also_saved_in_the_local_backup_folder(self):
        Task.objects.create(title="Keep a local copy", due_date=timezone.localdate())
        with tempfile.TemporaryDirectory() as directory:
            with override_settings(DATA_DIR=Path(directory)):
                response = self.client.get(reverse("backup"))
            backups = list((Path(directory) / "backups").glob("daily-backup-*.sqlite3"))
            self.assertEqual(len(backups), 1)
            self.assertEqual(backups[0].read_bytes(), response.content)

    def test_valid_backup_restores_task_data(self):
        Task.objects.create(title="Restore me", due_date=timezone.localdate())
        backup = self.client.get(reverse("backup"))
        Task.objects.all().delete()
        upload = SimpleUploadedFile("daily.sqlite3", backup.content, content_type="application/vnd.sqlite3")
        self.client.post(reverse("restore"), {"backup": upload})
        self.assertTrue(Task.objects.filter(title="Restore me").exists())

    def test_restoring_older_backup_migrates_schema_before_reporting_success(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "legacy.sqlite3"
            target = ("tasks", "0002_pomodorosession_pomo_completed_idx_and_more")
            legacy_connection = connections["default"]
            try:
                executor = MigrationExecutor(legacy_connection)
                executor.migrate([target])
                historical_apps = executor.loader.project_state([target]).apps
                old_task = historical_apps.get_model("tasks", "Task")
                old_task.objects.using("default").create(
                    title="Task from an older release", due_date=timezone.localdate()
                )
                legacy_connection.ensure_connection()
                with closing(sqlite3.connect(database_path)) as backup:
                    legacy_connection.connection.backup(backup)

                upload = SimpleUploadedFile(
                    "legacy.sqlite3", database_path.read_bytes(), content_type="application/vnd.sqlite3"
                )
                with override_settings(BACKUP_DIR=Path(directory) / "backups"):
                    response = self.client.post(reverse("restore"), {"backup": upload})

                self.assertEqual(response.status_code, 302)
                restored = Task.objects.get(title="Task from an older release")
                self.assertEqual(restored.category, "")
                current_executor = MigrationExecutor(connections["default"])
                self.assertFalse(
                    current_executor.migration_plan(current_executor.loader.graph.leaf_nodes())
                )
                self.assertTrue(list((Path(directory) / "backups").glob("pre-migration-*.sqlite3")))
            finally:
                legacy_connection.close()
                executor = MigrationExecutor(legacy_connection)
                executor.migrate(executor.loader.graph.leaf_nodes())

    def test_restore_rejects_unrecognized_database(self):
        data = io.BytesIO(b"not a database")
        data.name = "bad.sqlite3"
        response = self.client.post(reverse("restore"), {"backup": data})
        self.assertEqual(response.status_code, 302)
        self.assertContains(self.client.get(reverse("settings")), "Backup was not restored")


class ReminderSchedulerTests(TransactionTestCase):
    @patch("tasks.notifications.send_notification", return_value=True)
    def test_reminder_is_sent_once_and_expired_focus_is_completed(self, notify):
        now = timezone.now()
        task = Task.objects.create(title="Reminder", reminder_at=now - timedelta(minutes=1))
        session = PomodoroSession.objects.create(
            started_at=now - timedelta(minutes=26), duration_minutes=25
        )

        from .scheduler import check_reminders
        check_reminders()
        task.refresh_from_db()
        session.refresh_from_db()
        self.assertIsNotNone(task.notified_at)
        self.assertIsNotNone(session.completed_at)
        self.assertEqual(notify.call_count, 2)

        check_reminders()
        self.assertEqual(notify.call_count, 2)
